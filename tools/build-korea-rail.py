#!/usr/bin/env python3
"""Build the selected KTX corridors from Korail's public Oct 1 2026 workbook.
Usage: python3 tools/build-korea-rail.py [downloaded.xlsx]
Requires openpyxl. No API key. Routes are schematic station-to-station links.
"""
import datetime, json, pathlib, sys, urllib.request, tempfile
import openpyxl
ROOT=pathlib.Path(__file__).resolve().parent.parent
SOURCE='https://www.korail.com/file/cubedata/COMMON/jfile/202609/23/202609231a0cb38a2c5200.xlsx'
path=pathlib.Path(sys.argv[1]) if len(sys.argv)>1 else pathlib.Path(tempfile.gettempdir())/'ktx-20261001.xlsx'
if not path.exists(): urllib.request.urlretrieve(SOURCE,path)
# Approximate station anchors; geometry is explicitly a schematic, not track survey data.
anchors={
 '행신':[126.834,37.612], '서울':[126.972,37.555], '용산':[126.965,37.530], '광명':[126.884,37.416],
 '수서':[127.104,37.487], '동탄':[127.096,37.200], '평택지제':[127.070,37.018],
 '천안아산':[127.104,36.794], '오송':[127.328,36.620], '대전':[127.435,36.332],
 '김천구미':[128.181,36.113], '서대구':[128.540,35.881], '동대구':[128.629,35.880],
 '경주':[129.139,35.798], '울산':[129.138,35.551], '부산':[129.042,35.115],
 '공주':[127.104,36.332], '익산':[126.947,35.940], '정읍':[126.842,35.576],
 '광주송정':[126.792,35.138], '나주':[126.718,35.015], '목포':[126.386,34.791],
 '청량리':[127.047,37.580], '상봉':[127.086,37.596], '덕소':[127.209,37.586], '양평':[127.491,37.492],
 '서원주':[127.922,37.350], '만종':[127.921,37.354], '횡성':[128.011,37.483], '둔내':[128.222,37.511],
 '평창':[128.429,37.562], '진부(오대산)':[128.574,37.642], '강릉':[128.899,37.764],
}
# 서원주 and 만종 are separate stations, approximately 5 km apart.
anchors['서원주']=[127.837,37.350]
anchors['만종']=[127.893,37.354]
lines={
 'gyeongbu':{'name':'경부선','color':'#efb34f','tail':'오송 대전 김천구미 서대구 동대구 경주 울산 부산'.split()},
 'honam':{'name':'호남선','color':'#40c8be','tail':'오송 공주 익산 정읍 광주송정 나주 목포'.split()},
 'gangneung':{'name':'강릉선','color':'#ef8096','tail':'행신 서울 청량리 상봉 덕소 양평 서원주 만종 횡성 둔내 평창 진부(오대산) 강릉'.split()},
}
paths={}
for key,line in lines.items():
 if key=='gangneung':paths[key]=line['tail']
 else:
  paths[key+'-seoul']='행신 서울 용산 광명 천안아산'.split()+line['tail']
  paths[key+'-suseo']='수서 동탄 평택지제 천안아산'.split()+line['tail']
 del line['tail']
w=openpyxl.load_workbook(path,data_only=True)
trains=[];excluded=[];days='월화수목금토일'
for key,line in lines.items():
 s=w[line['name']];start=10 if key=='honam' else 11;header=start-3
 halves=[(2,4,19,20),(22,24,39,40)] if key=='gangneung' else [(2,4,24,25),(27,29,49,50)]
 for direction,(num,first,last,remarks) in enumerate(halves):
  for row in range(start,s.max_row+1):
   number=s.cell(row,num).value
   if not isinstance(number,(int,float)):continue
   stops=[];offset=0
   for col in range(first,last+1):
    value=s.cell(row,col).value
    if not isinstance(value,datetime.time) or value==datetime.time(0,0):continue
    minute=value.hour*60+value.minute+value.second/60
    name=str(s.cell(header,col).value).replace('\n','').replace(' ','')
    stops.append([name,minute])
   # Some reverse services are printed in the outbound half. Infer ordering from times.
   if len(stops)>2 and sum(b[1]<a[1] for a,b in zip(stops,stops[1:]))>1:stops.reverse()
   for i in range(1,len(stops)):
    while stops[i][1]<stops[i-1][1]:stops[i][1]+=1440
   route=key if key=='gangneung' else key+('-suseo' if any(x[0]=='수서' for x in stops) else '-seoul')
   if len(stops)<2 or any(x[0] not in paths[route] for x in stops):
    excluded.append({'line':key,'number':int(number),'row':row,'reason':'지원 범위 밖 경유역'});continue
   operating=str(s.cell(row,remarks).value).strip()
   assert operating=='매일' or all(x in days for x in operating),(key,row,operating)
   assert 0<stops[-1][1]-stops[0][1]<480,(key,row,stops)
   indexes=[paths[route].index(x[0]) for x in stops]
   direction=int(indexes[0]>indexes[-1])
   assert indexes==sorted(indexes,reverse=bool(direction)),(key,row,stops)
   trains.append({'id':f'{key}-{direction}-{int(number)}','number':int(number),'type':s.cell(row,num+1).value,'line':key,'route':route,'direction':direction,'days':list(range(7)) if operating=='매일' else [days.index(x) for x in operating],'stops':stops,'sourceRow':row})
result={'version':1,'effective':'2026-10-01','source':SOURCE,'sourceTitle':'코레일 KTX 시간표 (2026. 10. 1. 기준)','geometry':'Approximate station anchors and schematic links; not surveyed track geometry','positionMethod':'Constant speed between published station times; no dwell, delay or live location','scope':'경부·호남 고속선 및 강릉행. 수원·구포·서대전 경유, 동해행 및 기타 노선 제외. 요일별 정규 시간표; 임시운행 미반영.','stations':anchors,'lines':lines,'paths':paths,'trains':trains,'excluded':excluded}
out=ROOT/'data/korea.rail.json';out.write_text(json.dumps(result,ensure_ascii=False,separators=(',',':'))+'\n')
print(f'{len(trains)} timetable entries, {len(excluded)} excluded, {out.stat().st_size:,} bytes')
print({line:sum(t['line']==line for t in trains) for line in lines})
